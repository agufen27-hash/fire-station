"""Fase 15: mapeo exacto de celdas PCS/PCD2, escritura verificada contra la
plantilla en uso y PCD2 paginada de a 2 dotaciones con la firma validada."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from datetime import date, time

import openpyxl
import pytest
from PySide6.QtGui import QColor, QImage
from PySide6.QtWidgets import QApplication
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.db as db
import app.reports.excel_generator as eg
from app.core.catalogos import obtener_padron
from app.models import Personal
from app.services.personal_info import mandos_del_padron
from tests.ayudas import despachar, elegir_categoria


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def base_temporal(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}", future=True)
    monkeypatch.setattr(db, "engine", engine)
    monkeypatch.setattr(db, "SessionLocal", sessionmaker(bind=engine, autoflush=False, expire_on_commit=False))
    db.init_db()
    return engine


@pytest.fixture(scope="module")
def padron():
    return obtener_padron().bomberos


def _nombres():
    with db.get_session() as session:
        return {p.id_ruba: p.nombre_completo() for p in session.query(Personal).filter(Personal.id_ruba.isnot(None))}


# -- Escritura verificada ----------------------------------------------------------

def test_escritor_no_rompe_celdas_combinadas_ni_rotulos():
    wb = openpyxl.Workbook()
    hoja = wb.active
    hoja.merge_cells("C17:F17")
    hoja["C17"] = "Colision"
    hoja["A23"] = "Calle/Lugar:"
    hoja["I9"] = "/"
    e = eg.EscritorVerificado(hoja, "PCS")
    assert not e.escribir("D17", "Accidente - Tránsito", "Tipo")
    assert not e.escribir("A23", "PÉREZ, Juan", "Autorizó")
    assert e.escribir("I9", "23/09/2026", "Fecha")          # placeholder: se pisa
    assert hoja["C17"].value == "Colision" and hoja["A23"].value == "Calle/Lugar:"
    assert e.advertencias == [
        "PCS D17 (Tipo): está dentro del rango combinado C17:F17 ('Colision') — no se escribió.",
        "PCS A23 (Autorizó): la plantilla trae el rótulo 'Calle/Lugar:' — no se escribió.",
    ]
    aviso = eg.resumen_advertencias("PCS", eg.ResultadoPlanilla(None, e.advertencias))
    assert "2 celda(s)" in aviso and "D17, A23" in aviso


def test_dato_nunca_invisible_ni_cortado():
    """Descripción de la PCS rediseñada: letra blanca sobre relleno blanco."""
    from openpyxl.styles import Color, Font, PatternFill

    hoja = openpyxl.Workbook().active
    hoja["A26"].font = Font(color=Color(theme=0))
    hoja["A26"].fill = PatternFill("solid", fgColor=Color(theme=0))
    e = eg.EscritorVerificado(hoja, "PCS")
    assert e.escribir("A26", "Colisión en RN35", "Descripción")
    assert hoja["A26"].font.color.rgb == "FF000000"
    assert hoja["A26"].alignment.shrink_to_fit


# -- Servicio de prueba: 3 dotaciones + personal en base ----------------------------

def _firma_png(ruta):
    imagen = QImage(300, 90, QImage.Format.Format_ARGB32)
    imagen.fill(QColor("white"))
    for x in range(20, 280):
        imagen.setPixelColor(x, 45 + (x % 30) - 15, QColor("navy"))
    imagen.save(str(ruta))
    return ruta


@pytest.fixture
def servicio(qapp, base_temporal, padron, monkeypatch, tmp_path):
    """Accidente con 3 dotaciones (la 1 firmada con PIN) y personal en base
    (2 operadores + 14 en apresto, para pasar a la segunda columna de la PCS)."""
    from app.ui import main_window as mw

    monkeypatch.setattr(mw.QMessageBox, "warning", lambda *a, **k: pytest.fail(str(a[2:])))
    monkeypatch.setattr(mw, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(eg, "OUTPUT_DIR", tmp_path / "planillas")
    ventana = mw.MainWindow()

    jefe_1 = padron[2]
    with db.get_session() as session:
        persona = session.query(Personal).filter_by(id_ruba=jefe_1.id_ruba).one()
        persona.ruta_firma = str(_firma_png(tmp_path / "firma.png"))
        persona.pin = "4321"
        session.commit()

    elegir_categoria(ventana, mw, 1, _primer_codigo(ventana, mw, 1))
    ventana.entry_den_nombre.setText("Juan")
    ventana.entry_den_apellido.setText("Pérez")
    ventana.entry_den_dni.setText("30111222")
    ventana.entry_den_domicilio.setText("San Martín 450")
    ventana.combo_via_comunicacion.setCurrentText("WhatsApp")
    ventana.entry_contacto_detalle.setText("3584123456")
    ventana.selector_recibio.set_id_ruba(padron[30].id_ruba)
    ventana.radio_alarma_si.setChecked(True)
    mando = mandos_del_padron(padron)[0]
    ventana.selector_autorizo.set_id_ruba(mando.id_ruba)
    ventana.entry_calle.setText("Ruta 35 km 12")
    ventana.texto_resena.setPlainText("Colisión entre dos vehículos, sin atrapados.")
    ventana.campo_hora_llamado.set_value(time(14, 2))
    ventana.campo_hora_toque.set_value(time(14, 3))
    ventana.campo_fecha_salida.set_value(date(2026, 9, 23))
    ventana.campo_fecha_llegada.set_value(date(2026, 9, 23))
    ventana.campo_hora_salida.set_value(time(14, 5))
    ventana.campo_hora_llegada.set_value(time(15, 30))
    ventana._on_horario_general_cambiado()

    d1 = despachar(ventana, 4326, padron[0], (padron[1],), jefe=jefe_1)
    despachar(ventana, 9906, padron[3], (padron[4], padron[5]), jefe=padron[6])
    despachar(ventana, _otro_movil(ventana, {4326, 9906}), padron[7], (padron[8],), jefe=padron[9])
    assert d1.validar_pin("4321") is None and d1.esta_firmada()

    base = ventana.panel_base
    base.selector_operador_1.set_id_ruba(padron[10].id_ruba)
    base.selector_operador_2.set_id_ruba(padron[11].id_ruba)
    for bombero in padron[12:26]:                       # 14 en apresto
        base.agregar_apresto().set_id_ruba(bombero.id_ruba)

    resultado = ventana._persistir_incidente()
    assert resultado is not None
    yield resultado[0], resultado[1], padron, mando
    ventana.close()


def _primer_codigo(ventana, mw, tipo_id):
    ventana.combo_tipo.setCurrentIndex(ventana.combo_tipo.findData(tipo_id))
    for i in range(ventana.combo_categoria.count()):
        codigo = ventana.combo_categoria.itemData(i, mw.ROL_CODIGO_RUBA)
        if codigo:
            return codigo
    raise AssertionError("El tipo no tiene categorías")


def _otro_movil(ventana, usados):
    combo = ventana.panel_dotaciones.unidades()[0].combo_movil
    return next(combo.itemData(i) for i in range(combo.count()) if combo.itemData(i) not in (None, *usados))


def test_pcd2_pagina_de_a_dos_dotaciones_con_firma(servicio):
    incidente_id, numero, padron, _ = servicio
    resultado = eg.generar_planilla("PCD2", incidente_id)
    wb = openpyxl.load_workbook(resultado.ruta)
    nombre = _nombres()

    assert wb.sheetnames == ["Página 1", "Página 2"]
    p1, p2 = wb["Página 1"], wb["Página 2"]
    for n, hoja in enumerate((p1, p2), start=1):
        assert (hoja["C6"].value, hoja["S6"].value, hoja["V6"].value) == (numero, n, 2)
        assert hoja.print_area == "'Página %d'!$A$1:$W$53" % n
        assert hoja.sheet_properties.pageSetUpPr.fitToPage and hoja.page_setup.fitToHeight == 1  # 1 hoja A4
        assert len(hoja._images) >= 1                       # el logo también en la página 2

    assert (p1["A9"].value, p1["A32"].value, p2["A9"].value) == (1, 2, 3)   # numeración continua
    assert p2["A32"].value is None and p2["B34"].value is None
    assert (p1["B11"].value, p1["E14"].value, p1["E15"].value) == (
        nombre[padron[2].id_ruba], nombre[padron[0].id_ruba], nombre[padron[1].id_ruba])
    assert (p1["B34"].value, p1["E37"].value, p1["E38"].value, p1["E39"].value) == (
        nombre[padron[6].id_ruba], nombre[padron[3].id_ruba], nombre[padron[4].id_ruba], nombre[padron[5].id_ruba])
    assert (p2["B11"].value, p2["E14"].value, p2["E15"].value) == (
        nombre[padron[9].id_ruba], nombre[padron[7].id_ruba], nombre[padron[8].id_ruba])
    assert (p1["I9"].value, p1["L9"].value, p1["R9"].value, p1["U9"].value) == ("23/09/2026", "14:05", "23/09/2026", "15:30")
    assert p1["O11"].value                                   # grado del Jefe

    # Firma: solo la Dotación 1 validó el PIN -> imagen + aclaración en P26.
    assert "firma digital validada con PIN el" in p1["P26"].value
    assert nombre[padron[2].id_ruba] in p1["P26"].value
    assert p1["P49"].value is None and p2["P26"].value is None
    assert len(p1._images) == len(p2._images) + 1
    assert not [a for a in resultado.advertencias if "firma" in a]


def test_pcs_en_plantilla_rediseñada_completa_todas_las_celdas(servicio, tmp_path, monkeypatch):
    """Plantilla con las celdas del mapeo libres (como la rediseñada): se
    completa todo, sin advertencias."""
    incidente_id, numero, padron, mando = servicio
    carpeta = tmp_path / "plantillas"
    carpeta.mkdir()
    wb = openpyxl.Workbook()
    wb.save(carpeta / "PCS.xlsx")
    monkeypatch.setattr(eg, "TEMPLATES_DIR", carpeta)

    resultado = eg.generar_planilla("PCS", incidente_id)
    assert resultado.advertencias == []
    h = openpyxl.load_workbook(resultado.ruta).active
    nombre = _nombres()
    assert (h["C6"].value, h["J6"].value, h["T6"].value) == (numero, "Miércoles", "23/09/2026")
    assert (h["D9"].value, h["W9"].value, h["D10"].value) == ("Pérez, Juan", "14:02", "San Martín 450")
    assert (h["D11"].value, h["D12"].value) == ("WhatsApp: 3584123456", "30111222")
    assert (h["D13"].value, h["D14"].value, h["Q14"].value) == (nombre[padron[30].id_ruba], "SÍ", "14:03")
    assert h["D17"].value.startswith("Accidente - ")
    assert (h["D19"].value, h["D20"].value, h["D21"].value) == ("Adelia María", "Ruta 35 km 12", "Urbana")
    assert (h["A23"].value, h["A26"].value) == (nombre[mando.id_ruba], "Colisión entre dos vehículos, sin atrapados.")
    assert (h["F31"].value, h["F32"].value) == (nombre[padron[10].id_ruba], nombre[padron[11].id_ruba])
    apresto = [nombre[b.id_ruba] for b in padron[12:26]]
    assert [h[f"B{f}"].value for f in range(35, 48)] == apresto[:13]
    assert (h["O35"].value, h["O36"].value) == (apresto[13], None)


def test_pcs_en_plantilla_actual_avisa_sin_romperla(servicio):
    """Con la plantilla vigente (todavía sin rediseñar) las celdas ocupadas
    por rótulos/combinados se saltean con aviso; las libres se completan."""
    incidente_id, _, padron, _ = servicio
    plantilla = openpyxl.load_workbook(eg._ubicar_plantilla("PCS.xlsx")).active
    resultado = eg.generar_planilla("PCS", incidente_id)
    h = openpyxl.load_workbook(resultado.ruta).active

    libres = [c for c in ("D12", "D13", "F32") if eg.EscritorVerificado(plantilla, "PCS")._problema(c) is None]
    for celda in libres:
        assert h[celda].value, celda
    ocupadas = {c for c in ("D11", "D17", "A23") if eg.EscritorVerificado(plantilla, "PCS")._problema(c)}
    for celda in ocupadas:
        assert h[celda].value == plantilla[celda].value    # la plantilla queda intacta
        assert any(a.startswith(f"PCS {celda} ") for a in resultado.advertencias)
