"""
Importación de los reportes Excel que exporta RUBA a la base local
(data/fire_station.db). El archivo SIEMPRE lo elige el usuario (QFileDialog
en la UI): este módulo no conoce ninguna ruta fija y no depende de Qt.

- 'Reporte de vehiculos.xlsx'  -> tabla `moviles`  (clave: 'Nº Móvil').
- 'Reporte de bomberos.xlsx'   -> tabla `personal` (clave: 'Id' de RUBA, luego DNI).

Cada importación corre en UNA transacción (get_session): si una fila rompe
una restricción de la base, no queda nada a medio importar.

Particularidades del export real de RUBA que se contemplan:
  - 'Nº Móvil' viene a veces como número (26) y con espacios dobles
    ('Rojo  22'); en la base el móvil puede figurar como 'Móvil 26'.
    Todo se compara por una clave normalizada (ver `clave_movil`).
  - El mismo móvil puede aparecer dos veces (un alta vieja y la vigente):
    se queda la fila más completa y se avisa en el resumen.
  - El reporte de bomberos real NO trae 'Jerarquía': se usa esa columna si
    existe y, si no, 'Formación' (Bombero / Aspirante / Cadete).
  - Personal: lectura híbrida por nombre de encabezado + letra de columna
    (ver CAMPOS_PADRON), fechas normalizadas con `parsear_fecha_ruba` y
    upsert que nunca toca el PIN de un bombero existente.
"""

from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Sequence, Tuple

from sqlalchemy.orm import Session

from app.core.app_log import asegurar_log_app

ESTADO_MOVIL_EN_SERVICIO = "En Servicio"
ESTADO_MOVIL_FUERA_DE_SERVICIO = "Fuera de Servicio"
ESTADO_MOVIL_BAJA = "Baja"

ESTADOS_PERSONAL = ("Activo", "Licencia", "Reserva", "Baja")


# ---------------------------------------------------------------------------
# Lectura genérica de la hoja
# ---------------------------------------------------------------------------

def _clave(texto: Any) -> str:
    """Sin acentos, mayúsculas, espacios colapsados (para encabezados y valores)."""
    descompuesto = unicodedata.normalize("NFKD", str(texto or ""))
    sin_acentos = "".join(c for c in descompuesto if not unicodedata.combining(c))
    return " ".join(sin_acentos.upper().split())


def _texto(valor: Any) -> Optional[str]:
    """Celda -> texto limpio (espacios colapsados); 14000.0 -> '14000'; None si vacía."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = " ".join(str(valor).split())
    return texto or None


def _es_si(valor: Any) -> bool:
    return _clave(valor) in {"SI", "S", "TRUE", "1", "VERDADERO"}


def _entero(valor: Any) -> Optional[int]:
    try:
        return int(float(str(valor).strip())) if valor not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _fecha(valor: Any) -> Optional[date]:
    """RUBA exporta dd/mm/aa como texto; openpyxl puede dar datetime."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto or "", formato).date()
        except ValueError:
            continue
    return None


def _leer_hoja(ruta: Path, requeridas: List[str]) -> Tuple[Dict[str, int], Iterator[tuple]]:
    """Primera hoja del Excel: (índice de columnas por clave normalizada, filas).
    FileNotFoundError si no existe; ValueError si no es .xlsx o faltan columnas."""
    from zipfile import BadZipFile

    from openpyxl import load_workbook  # import diferido: openpyxl es pesado
    from openpyxl.utils.exceptions import InvalidFileException

    ruta = Path(ruta)
    if not ruta.is_file():
        raise FileNotFoundError(f"No existe el archivo: {ruta}")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError) as e:
        raise ValueError(f"{ruta.name} no es un Excel (.xlsx) válido: {e}") from e
    try:
        filas = list(libro.worksheets[0].iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        raise ValueError(f"{ruta.name}: la hoja está vacía.")
    indices = {_clave(h): i for i, h in enumerate(filas[0]) if h is not None}
    faltantes = [c for c in requeridas if _clave(c) not in indices]
    if faltantes:
        raise ValueError(f"{ruta.name}: faltan las columnas {faltantes}. ¿Es el reporte correcto de RUBA?")
    return indices, iter(filas[1:])


def _celda(indices: Dict[str, int], fila: tuple, columna: str) -> Any:
    i = indices.get(_clave(columna))
    return fila[i] if i is not None and i < len(fila) else None


@dataclass
class ResumenImportacion:
    total: int = 0
    nuevas: int = 0
    actualizadas: int = 0
    avisos: List[str] = field(default_factory=list)
    ignoradas: int = 0

    def texto(self, sustantivo: str) -> str:
        return f"{self.total} {sustantivo} importadas ({self.nuevas} nuevas, {self.actualizadas} actualizadas)"

    def texto_personal(self) -> str:
        return (f"Importación completada: {self.actualizadas} bomberos actualizados, "
                f"{self.nuevas} dados de alta, {self.ignoradas} filas ignoradas")


# ---------------------------------------------------------------------------
# Vehículos
# ---------------------------------------------------------------------------

COLUMNAS_VEHICULOS = ["Nº Móvil", "Marca", "Modelo", "Tipo", "Año", "Servicio", "Baja"]

_RE_PREFIJO_MOVIL = re.compile(r"^(MOVIL|MOV\.?|N[º°O]?\s*MOVIL:?)\s+(?=\d+$)")


def normalizar_dominio(valor: Any) -> Optional[str]:
    """'ab 123 cd' / 'AB-123-CD' -> 'AB123CD'. None si queda vacío."""
    texto = re.sub(r"[\s.\-]", "", _texto(valor) or "").upper()
    return texto[:15] or None


def clave_movil(numero: Any) -> str:
    """Clave de comparación de un móvil: 'Rojo  22' == 'ROJO 22';
    26 == '26' == 'Móvil 26' (así calza con los nombres ya cargados)."""
    return _RE_PREFIJO_MOVIL.sub("", _clave(_texto(numero)))


def nombre_para_movil(numero: str) -> str:
    """Nombre visible de una unidad nueva: mismo criterio que el alta desde
    ruba_mapping.json (app/db.py): '26' -> 'Móvil 26', 'Rojo 18' tal cual."""
    return f"Móvil {numero}" if numero.isdigit() else numero


@dataclass
class VehiculoRuba:
    numero: str                        # 'Nº Móvil' limpio
    tipo: Optional[str]
    marca: Optional[str]
    modelo: Optional[str]
    anio: Optional[int]
    estado: str                        # ESTADO_MOVIL_*
    fecha_adquisicion: Optional[date] = None
    dominio: Optional[str] = None      # solo si el reporte trae "Dominio" / "Patente"
    fila_excel: int = 0
    completitud: int = 0               # para elegir entre filas duplicadas

    @property
    def activo(self) -> bool:
        return self.estado == ESTADO_MOVIL_EN_SERVICIO

    @property
    def clave(self) -> str:
        return clave_movil(self.numero)


def _estado_vehiculo(servicio: Any, baja: Any) -> str:
    if _es_si(baja):
        return ESTADO_MOVIL_BAJA
    if _es_si(servicio):
        return ESTADO_MOVIL_EN_SERVICIO
    return ESTADO_MOVIL_FUERA_DE_SERVICIO


def _preferir(a: VehiculoRuba, b: VehiculoRuba) -> VehiculoRuba:
    """Entre dos filas del mismo móvil: la que no está de baja, luego la más
    completa, luego la adquirida más recientemente, luego la de más abajo."""
    def orden(v: VehiculoRuba):
        return (v.estado != ESTADO_MOVIL_BAJA, v.completitud, v.fecha_adquisicion or date.min, v.fila_excel)
    return max(a, b, key=orden)


def leer_reporte_vehiculos(ruta: Path) -> Tuple[List[VehiculoRuba], List[str]]:
    """Filas del 'Reporte de vehiculos' de RUBA, sin duplicados por Nº Móvil.
    Devuelve (vehículos, avisos)."""
    indices, filas = _leer_hoja(ruta, COLUMNAS_VEHICULOS)
    avisos: List[str] = []
    por_clave: Dict[str, VehiculoRuba] = {}
    for n, fila in enumerate(filas, start=2):
        numero = _texto(_celda(indices, fila, "Nº Móvil"))
        if not numero:
            continue  # fila vacía / totales
        marca = _texto(_celda(indices, fila, "Marca")) or _texto(_celda(indices, fila, "Otra Marca"))
        modelo = _texto(_celda(indices, fila, "Modelo"))
        tipo = _texto(_celda(indices, fila, "Tipo"))
        vto = _texto(_celda(indices, fila, "Fecha de Vto. Verificacion"))
        vehiculo = VehiculoRuba(
            numero=numero, tipo=tipo, marca=marca, modelo=modelo,
            anio=_entero(_celda(indices, fila, "Año")),
            estado=_estado_vehiculo(_celda(indices, fila, "Servicio"), _celda(indices, fila, "Baja")),
            fecha_adquisicion=_fecha(_celda(indices, fila, "Fecha de Adquisición")),
            dominio=normalizar_dominio(_celda(indices, fila, "Dominio") or _celda(indices, fila, "Patente")),
            fila_excel=n,
            # un modelo puramente numérico igual al año suele ser un alta mal cargada
            completitud=sum(1 for x in (marca, tipo, vto) if x) + (1 if modelo and not modelo.isdigit() else 0),
        )
        previo = por_clave.get(vehiculo.clave)
        if previo is None:
            por_clave[vehiculo.clave] = vehiculo
            continue
        elegido = _preferir(previo, vehiculo)
        descartado = vehiculo if elegido is previo else previo
        por_clave[vehiculo.clave] = elegido
        avisos.append(
            f"'{descartado.numero}' (fila {descartado.fila_excel}) repite el móvil '{elegido.numero}' "
            f"(fila {elegido.fila_excel}): se tomó la fila {elegido.fila_excel} "
            f"({elegido.marca or 's/marca'} {elegido.modelo or ''}, {elegido.estado})."
        )
    if not por_clave:
        raise ValueError(f"{Path(ruta).name}: no se encontró ningún móvil.")
    return list(por_clave.values()), avisos


def _ids_ruba_por_clave() -> Dict[str, int]:
    """Nº Móvil -> value del <option> de RUBA (ruba_mapping.json), si se conoce."""
    try:
        from app.core.catalogos import cargar_moviles

        return {clave_movil(m.numero): m.id_ruba for m in cargar_moviles()}
    except (OSError, ValueError):
        return {}


def upsert_vehiculos(session: Session, vehiculos: List[VehiculoRuba]) -> ResumenImportacion:
    """Alta o actualización de cada vehículo usando 'Nº Móvil' como clave.
    Se busca primero por `numero_movil` (importaciones previas) y después
    por el nombre visible ('Rojo 18', 'Móvil 26'), para enlazar las
    unidades ya cargadas en vez de duplicarlas. Nunca borra unidades."""
    from app.models import Movil

    resumen = ResumenImportacion(total=len(vehiculos))
    ids_ruba = _ids_ruba_por_clave()
    existentes = session.query(Movil).all()
    por_numero = {clave_movil(m.numero_movil): m for m in existentes if m.numero_movil}
    por_nombre = {clave_movil(m.nombre_identificador): m for m in existentes}
    nombres_usados = {m.nombre_identificador for m in existentes}

    for v in vehiculos:
        movil = por_numero.get(v.clave) or por_nombre.get(v.clave)
        if movil is None:
            nombre = nombre_para_movil(v.numero)
            if nombre in nombres_usados:  # choca con otra unidad cargada a mano
                nombre = f"{nombre} (RUBA)"
            movil = Movil(nombre_identificador=nombre)
            session.add(movil)
            nombres_usados.add(nombre)
            resumen.nuevas += 1
        else:
            resumen.actualizadas += 1
        movil.numero_movil = v.numero
        movil.tipo = v.tipo
        movil.marca = v.marca
        movil.modelo = v.modelo
        movil.anio = v.anio
        if v.dominio:  # el reporte habitual no trae dominio: no se borra el cargado a mano
            movil.dominio = v.dominio
        movil.estado = v.estado
        movil.activo = v.activo
        if movil.id_ruba is None and v.clave in ids_ruba:
            otro = session.query(Movil).filter(Movil.id_ruba == ids_ruba[v.clave]).first()
            if otro is None or otro is movil:
                movil.id_ruba = ids_ruba[v.clave]
        por_numero[v.clave] = movil
    session.flush()
    return resumen


def importar_vehiculos_desde_excel(ruta: Path) -> ResumenImportacion:
    """Lee y aplica el 'Reporte de vehiculos' en una sola transacción."""
    from app.db import get_session

    vehiculos, avisos = leer_reporte_vehiculos(ruta)
    with get_session() as session:
        resumen = upsert_vehiculos(session, vehiculos)
    resumen.avisos = avisos + resumen.avisos
    return resumen


# ---------------------------------------------------------------------------
# Bomberos
# ---------------------------------------------------------------------------
#
# Lectura HÍBRIDA del 'Reporte de bomberos': cada campo tiene una letra de
# columna de referencia y los nombres de encabezado con que RUBA lo exporta.
# Se busca primero por NOMBRE normalizado (minúsculas, sin tildes, sin
# espacios ni signos: "Jerarquía " == "jerarquia") y, si no aparece, se usa
# la LETRA solo si el encabezado de esa columna está vacío o coincide. Así:
#   - si RUBA agrega o quita columnas intermedias, el nombre lo encuentra;
#   - una letra que hoy trae otro dato (en el export actual W es
#     "Formación", no "Fecha Ingreso") nunca se lee como si fuera otro campo.

log = logging.getLogger(__name__)

# (campo, letra de referencia, encabezados aceptados)
CAMPOS_PADRON: Tuple[Tuple[str, Optional[str], Tuple[str, ...]], ...] = (
    ("id_ruba", "A", ("Id", "ID RUBA", "Id Sistema")),
    ("activo", "B", ("Activo",)),
    ("apellido", "C", ("Apellido", "Apellidos")),
    ("nombre", "D", ("Nombre", "Nombres")),
    ("dni", "E", ("DNI", "Documento", "Nro Documento", "N° Documento")),
    ("fecha_nacimiento", "F", ("Fecha de nacimiento", "Fecha nacimiento", "Nacimiento")),
    ("legajo", "G", ("Legajo", "N° Legajo", "Nro Legajo")),
    ("sexo", "J", ("Sexo", "Género", "Genero")),
    ("domicilio", "M", ("Domicilio", "Dirección")),
    ("email", "N", ("Email", "E-mail", "Correo", "Correo electrónico")),
    ("telefono", "O", ("Telefono", "Teléfono", "Celular")),
    ("grupo_sanguineo", "P", ("Grupo Sanguineo", "Grupo Sanguíneo", "Grupo")),
    ("factor", "Q", ("Factor", "Factor RH", "RH")),
    ("fecha_ingreso", "W", ("Fecha Ingreso", "Fecha de Ingreso", "Ingreso")),
    ("cargo", "Y", ("Cargo",)),
    ("jerarquia", "AA", ("Jerarquía", "Jerarquia", "Grado")),
    ("nivel_educativo", "AF", ("Nivel Educativo",)),
    ("titulo_obtenido", "AG", ("Título obtenido", "Titulo obtenido", "Título", "Titulo")),
    ("estado", "AI", ("Estado", "Situación de revista", "Situacion de revista", "Revista")),
    ("fecha_ultimo_ascenso", "AL", ("Fecha último Ascenso", "Fecha ultimo ascenso", "Fecha de último ascenso",
                                    "Último ascenso")),
    # Sin letra fija: solo por nombre (completa la jerarquía si no hay columna 'Jerarquía').
    ("formacion", None, ("Formación", "Formacion")),
)
CAMPOS_FECHA = ("fecha_nacimiento", "fecha_ingreso", "fecha_ultimo_ascenso")
MARCAS_HEADER = {"legajo", "apellido", "dni"}
FILAS_BUSCAR_HEADER = 10

_VACIOS = {"", "-", "--", "S/D", "SD", "N/A", "NA", "SIN DATO", "NULL", "NONE"}
_EPOCA_EXCEL = date(1899, 12, 30)
_FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%Y/%m/%d", "%d.%m.%Y")
_SI_ACTIVO = {"SI", "S", "1", "TRUE", "VERDADERO", "ACTIVO", "ACTIVA", "X"}


def _norm_header(texto: Any) -> str:
    """'Jerarquía ' -> 'jerarquia'; 'Fecha de nacimiento' -> 'fechadenacimiento'."""
    descompuesto = unicodedata.normalize("NFKD", str(texto or ""))
    sin_acentos = "".join(c for c in descompuesto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", sin_acentos.strip().lower())


def parsear_fecha_ruba(valor: Any) -> Optional[str]:
    """Cualquier fecha del export de RUBA -> 'YYYY-MM-DD' (ISO), o None si
    está vacía o no se entiende. Nunca lanza.
      - datetime / date -> su fecha.
      - int / float -> número de serie de Excel (época 1899-12-30).
      - texto -> sin espacios y sin la hora pegada ('2009-07-08 00:00:00'),
        en dd/mm/aaaa, aaaa-mm-dd, dd-mm-aaaa (y dd/mm/aa).
      - '', '-', 'S/D'... -> None."""
    try:
        if valor is None or isinstance(valor, bool):
            return None
        if isinstance(valor, datetime):
            return valor.date().isoformat()
        if isinstance(valor, date):
            return valor.isoformat()
        if isinstance(valor, (int, float)):
            serial = float(valor)
            if not 1 <= serial <= 2_958_465:  # 31/12/9999 en Excel
                return None
            return (_EPOCA_EXCEL + timedelta(days=int(serial))).isoformat()
        texto = str(valor).strip()
        if texto.upper() in _VACIOS:
            return None
        texto = re.split(r"[ T]", texto, maxsplit=1)[0]  # sin la hora adosada
        if re.fullmatch(r"\d+(\.\d+)?", texto):           # serial de Excel guardado como texto
            return parsear_fecha_ruba(float(texto))
        for formato in _FORMATOS_FECHA:
            try:
                return datetime.strptime(texto, formato).date().isoformat()
            except ValueError:
                continue
        return None
    except (TypeError, ValueError, OverflowError):
        return None


def _a_date(iso: Optional[str]) -> Optional[date]:
    return date.fromisoformat(iso) if iso else None


def _texto_limpio(valor: Any) -> Optional[str]:
    """strip + espacios colapsados; marcas de vacío ('-', 'S/D') -> None."""
    texto = _texto(valor)
    return None if texto is None or texto.upper() in _VACIOS else texto


def _activo_ruba(valor: Any) -> bool:
    """'SI', 'S', '1', 'TRUE', 'ACTIVO' -> True; 'NO', '0', 'FALSE', vacío -> False."""
    if isinstance(valor, bool):
        return valor
    return _clave(_texto(valor)) in _SI_ACTIVO


def _solo_digitos(valor: Any) -> Optional[str]:
    """DNI: '24.783.126' / 24783126.0 / '24 783 126' -> '24783126'."""
    texto = re.sub(r"\D", "", _texto(valor) or "")
    return texto or None


def _legajo_limpio(valor: Any) -> Optional[str]:
    """Legajo sin puntos de miles, comas ni espacios. Se respeta el guion de
    los legajos del cuerpo ('59-130'); '/' se unifica a '-'."""
    texto = _texto_limpio(valor)
    if not texto:
        return None
    texto = re.sub(r"[.,\s]", "", texto).replace("/", "-")
    return texto or None


def _grupo_sanguineo(grupo: Optional[str], factor: Optional[str]) -> Optional[str]:
    """'A' + 'Positivo' -> 'A+' (cabe en la columna y es como se lee en la ficha)."""
    if not grupo:
        return None
    grupo = grupo.upper().replace(" ", "")
    if grupo.endswith(("+", "-")) or not factor:
        return grupo[:10]
    signo = "+" if _clave(factor).startswith(("POS", "+")) else "-" if _clave(factor).startswith(("NEG", "-")) else ""
    return f"{grupo}{signo}"[:10]


@dataclass
class BomberoRuba:
    id_ruba: Optional[int]
    apellido: str
    nombre: str
    dni: Optional[str]
    legajo: Optional[str]
    jerarquia: Optional[str]
    jerarquia_es_oficial: bool          # vino de 'Jerarquía' (no de 'Formación')
    estado: str                         # uno de ESTADOS_PERSONAL
    activo: bool = True                 # columna 'Activo' (B)
    fila: int = 0                       # fila del Excel (para los avisos)
    fecha_nacimiento: Optional[date] = None
    sexo: Optional[str] = None
    domicilio: Optional[str] = None
    email: Optional[str] = None
    telefono: Optional[str] = None
    grupo_sanguineo: Optional[str] = None
    factor_rh: Optional[str] = None
    fecha_ingreso: Optional[date] = None
    cargo: Optional[str] = None
    formacion: Optional[str] = None
    nivel_educativo: Optional[str] = None
    titulo_obtenido: Optional[str] = None
    situacion_revista: Optional[str] = None
    fecha_ultimo_ascenso: Optional[date] = None

    @property
    def descripcion(self) -> str:
        nombre = ", ".join(x for x in (self.apellido, self.nombre) if x) or "(sin nombre)"
        return f"fila {self.fila}: {nombre}" + (f" (Id RUBA {self.id_ruba})" if self.id_ruba else "")


@dataclass
class LecturaPadron:
    bomberos: List[BomberoRuba]
    ignoradas: int = 0
    avisos: List[str] = field(default_factory=list)
    columnas: Dict[str, str] = field(default_factory=dict)   # campo -> letra realmente usada


def _estado_personal(estado: Any, activo: Any) -> str:
    clave = _clave(estado)
    for opcion in ESTADOS_PERSONAL:
        if clave == _clave(opcion):
            return opcion
    if clave.startswith("LICENCIA"):
        return "Licencia"
    if clave:  # 'Suspendido', 'Inactivo'... cualquier otro estado de RUBA
        return "Baja" if clave not in {"ACTIVO", "ACTIVA"} else "Activo"
    return "Activo" if activo is None or _activo_ruba(activo) else "Baja"


def _buscar_fila_encabezado(filas: List[tuple]) -> int:
    """Índice de la fila de encabezados: la primera (de las primeras 10) que
    tiene al menos dos de 'Legajo', 'Apellido', 'DNI'."""
    for i, fila in enumerate(filas[:FILAS_BUSCAR_HEADER]):
        normalizados = {_norm_header(v) for v in fila if v is not None}
        if len(MARCAS_HEADER & normalizados) >= 2:
            return i
    raise ValueError("No se encontró la fila de encabezados (Legajo / Apellido / DNI) en las primeras "
                     f"{FILAS_BUSCAR_HEADER} filas. ¿Es el 'Reporte de bomberos' de RUBA?")


def resolver_columnas(encabezado: Sequence[Any]) -> Tuple[Dict[str, int], List[str]]:
    """{campo: índice de columna} con la estrategia híbrida + avisos."""
    from openpyxl.utils import column_index_from_string

    por_nombre: Dict[str, int] = {}
    for i, valor in enumerate(encabezado):
        clave = _norm_header(valor)
        if clave and clave not in por_nombre:
            por_nombre[clave] = i
    indices: Dict[str, int] = {}
    avisos: List[str] = []
    for campo, letra, nombres in CAMPOS_PADRON:
        aceptados = {_norm_header(n) for n in nombres}
        encontrado = next((por_nombre[n] for n in (_norm_header(x) for x in nombres) if n in por_nombre), None)
        if encontrado is None and letra:
            i = column_index_from_string(letra) - 1
            actual = encabezado[i] if i < len(encabezado) else None
            if i < len(encabezado) and (actual in (None, "") or _norm_header(actual) in aceptados):
                encontrado = i
            elif actual not in (None, ""):
                avisos.append(f"'{nombres[0]}': la columna {letra} trae '{str(actual).strip()}' y no hay otra "
                              f"con ese encabezado; el dato no se importa.")
        if encontrado is not None:
            indices[campo] = encontrado
    return indices, avisos


def leer_reporte_bomberos(ruta: Path) -> LecturaPadron:
    """TODAS las filas del 'Reporte de bomberos' (Activo, Reserva, Baja...),
    a diferencia de catalogos.leer_padron_personal, que solo deja al
    personal activo para los selectores de RUBA. Las filas incompletas se
    informan y se saltean sin cortar la lectura."""
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException

    asegurar_log_app()
    ruta = Path(ruta)
    if not ruta.is_file():
        raise FileNotFoundError(f"No existe el archivo: {ruta}")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError) as e:
        raise ValueError(f"{ruta.name} no es un Excel (.xlsx) válido: {e}") from e
    try:
        filas = list(libro.worksheets[0].iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        raise ValueError(f"{ruta.name}: la hoja está vacía.")

    fila_header = _buscar_fila_encabezado(filas)
    indices, avisos = resolver_columnas(filas[fila_header])
    faltan = [c for c in ("apellido", "nombre") if c not in indices]
    if faltan or ("id_ruba" not in indices and "dni" not in indices):
        raise ValueError(f"{ruta.name}: faltan columnas obligatorias ({', '.join(faltan) or 'Id o DNI'}). "
                         "¿Es el 'Reporte de bomberos' de RUBA?")
    lectura = LecturaPadron([], avisos=list(avisos),
                            columnas={c: get_column_letter(i + 1) for c, i in indices.items()})
    log.info("Padrón %s: encabezados en la fila %s; columnas %s", ruta.name, fila_header + 1, lectura.columnas)
    for aviso in avisos:
        log.warning("Padrón %s: %s", ruta.name, aviso)

    def celda(fila: tuple, campo: str) -> Any:
        i = indices.get(campo)
        return fila[i] if i is not None and i < len(fila) else None

    def fecha(fila: tuple, campo: str, b_desc: str) -> Optional[date]:
        crudo = celda(fila, campo)
        iso = parsear_fecha_ruba(crudo)
        if iso is None and _texto_limpio(crudo) and not isinstance(crudo, bool):
            log.warning("Padrón: %s: '%s' no es una fecha válida (%s); queda vacía.", b_desc, crudo, campo)
        return _a_date(iso)

    hay_jerarquia = "jerarquia" in indices
    for numero, fila in enumerate(filas[fila_header + 1:], start=fila_header + 2):
        if not any(v not in (None, "") for v in fila):
            continue  # fila vacía: ni se cuenta
        id_ruba = _entero(celda(fila, "id_ruba"))
        dni = _solo_digitos(celda(fila, "dni"))
        apellido = (_texto_limpio(celda(fila, "apellido")) or "").upper()
        nombre = _texto_limpio(celda(fila, "nombre")) or ""
        if id_ruba is None and not dni:
            lectura.ignoradas += 1
            log.warning("Padrón: fila %s ignorada: no tiene Id de RUBA ni DNI (%s).", numero,
                        ", ".join(x for x in (apellido, nombre) if x) or "sin nombre")
            continue
        if not apellido and not nombre:
            lectura.ignoradas += 1
            log.warning("Padrón: fila %s ignorada: sin apellido ni nombre (Id %s).", numero, id_ruba)
            continue
        activo = _activo_ruba(celda(fila, "activo")) if "activo" in indices else True
        situacion = _texto_limpio(celda(fila, "estado"))
        formacion = _texto_limpio(celda(fila, "formacion"))
        jerarquia = _texto_limpio(celda(fila, "jerarquia")) if hay_jerarquia else formacion
        factor = _texto_limpio(celda(fila, "factor"))
        b = BomberoRuba(
            id_ruba=id_ruba, apellido=apellido, nombre=nombre, dni=dni,
            legajo=_legajo_limpio(celda(fila, "legajo")),
            jerarquia=jerarquia, jerarquia_es_oficial=hay_jerarquia,
            estado=_estado_personal(situacion, activo if "activo" in indices else None),
            activo=activo, fila=numero,
            sexo=_texto_limpio(celda(fila, "sexo")),
            domicilio=_texto_limpio(celda(fila, "domicilio")),
            email=(_texto_limpio(celda(fila, "email")) or "").lower() or None,
            telefono=_texto_limpio(celda(fila, "telefono")),
            grupo_sanguineo=_grupo_sanguineo(_texto_limpio(celda(fila, "grupo_sanguineo")), factor),
            factor_rh=factor,
            cargo=_texto_limpio(celda(fila, "cargo")),
            formacion=formacion,
            nivel_educativo=_texto_limpio(celda(fila, "nivel_educativo")),
            titulo_obtenido=_texto_limpio(celda(fila, "titulo_obtenido")),
            situacion_revista=situacion,
        )
        b.fecha_nacimiento = fecha(fila, "fecha_nacimiento", b.descripcion)
        b.fecha_ingreso = fecha(fila, "fecha_ingreso", b.descripcion)
        b.fecha_ultimo_ascenso = fecha(fila, "fecha_ultimo_ascenso", b.descripcion)
        incompletos = [n for n, v in (("DNI", dni), ("Legajo", b.legajo), ("Id RUBA", id_ruba)) if not v]
        if incompletos:
            log.warning("Padrón: %s sin %s.", b.descripcion, ", ".join(incompletos))
        lectura.bomberos.append(b)
    if not lectura.bomberos:
        raise ValueError(f"{ruta.name}: no se encontró ningún bombero con Id o DNI.")
    return lectura


# Campos que la importación copia a `Personal` SOLO si el Excel trae el dato
# (un vacío en RUBA nunca borra lo que se cargó a mano en Fire Station).
_CAMPOS_COMPLETAR = (
    "fecha_nacimiento", "sexo", "domicilio", "email", "telefono", "grupo_sanguineo", "factor_rh",
    "cargo", "formacion", "nivel_educativo", "titulo_obtenido", "situacion_revista", "fecha_ultimo_ascenso",
)


def upsert_bomberos(session: Session, bomberos: List[BomberoRuba]) -> ResumenImportacion:
    """UPSERT por Id de RUBA, luego DNI y luego Legajo (así el personal ya
    cargado conserva su legajo digital, firma y PIN). Reglas:
      - PIN: JAMÁS se toca el de un bombero existente; solo un alta nueva
        recibe el PIN de fábrica (hasheado) y se le pide cambiarlo al usarlo.
      - estado/activo, apellido, nombre e id_ruba se toman de RUBA;
      - legajo y DNI se actualizan solo si no los usa otra persona;
      - los demás datos (fechas, contacto, cargo...) solo se escriben si el
        Excel los trae: un vacío no borra un dato cargado a mano;
      - la jerarquía se pisa solo si viene de la columna 'Jerarquía'; la
        'Formación' (Bombero/Aspirante) solo completa jerarquías vacías;
      - a quien está de Baja y no existe en la base no se lo da de alta."""
    from app.core.security import PIN_DEFAULT, hash_pin
    from app.models import Personal

    resumen = ResumenImportacion()
    existentes = session.query(Personal).all()
    por_id = {p.id_ruba: p for p in existentes if p.id_ruba is not None}
    por_dni = {p.dni: p for p in existentes}
    legajos = {p.legajo: p for p in existentes if p.legajo}

    for b in bomberos:
        persona = ((por_id.get(b.id_ruba) if b.id_ruba is not None else None)
                   or (por_dni.get(b.dni) if b.dni else None)
                   or (legajos.get(b.legajo) if b.legajo else None))
        if persona is None:
            if b.estado == "Baja":
                resumen.ignoradas += 1
                log.info("Padrón: %s está de Baja y no existe en la base: no se da de alta.", b.descripcion)
                continue
            if not b.dni:
                resumen.ignoradas += 1
                aviso = f"{b.descripcion} no tiene DNI: no se dio de alta (completalo en RUBA o cargalo a mano)."
                resumen.avisos.append(aviso)
                log.warning("Padrón: %s", aviso)
                continue
            # Alta nueva: único caso en que se fija el PIN (el de fábrica, hasheado).
            persona = Personal(dni=b.dni, nombre=b.nombre, apellido=b.apellido, pin=hash_pin(PIN_DEFAULT))
            session.add(persona)
            por_dni[b.dni] = persona
            resumen.nuevas += 1
        else:
            resumen.actualizadas += 1
        resumen.total += 1

        if b.id_ruba is not None:
            persona.id_ruba = b.id_ruba
            por_id[b.id_ruba] = persona
        persona.apellido = b.apellido or persona.apellido
        persona.nombre = b.nombre or persona.nombre
        if b.dni and persona.dni != b.dni:
            if por_dni.get(b.dni) not in (None, persona):
                aviso = f"DNI {b.dni} de {b.apellido}, {b.nombre} ya es de otra persona: no se cambió."
                resumen.avisos.append(aviso)
                log.warning("Padrón: %s", aviso)
            else:
                por_dni.pop(persona.dni, None)
                persona.dni = b.dni
                por_dni[b.dni] = persona
        if b.legajo and persona.legajo != b.legajo:
            if legajos.get(b.legajo) not in (None, persona):
                aviso = f"Legajo {b.legajo} de {b.apellido}, {b.nombre} ya está asignado: no se cambió."
                resumen.avisos.append(aviso)
                log.warning("Padrón: %s", aviso)
            else:
                legajos.pop(persona.legajo, None)
                persona.legajo = b.legajo
                legajos[b.legajo] = persona
        if b.jerarquia and (b.jerarquia_es_oficial or not persona.jerarquia):
            persona.jerarquia = b.jerarquia
        if b.fecha_ingreso:
            persona.antiguedad_fecha = b.fecha_ingreso
        for campo in _CAMPOS_COMPLETAR:
            valor = getattr(b, campo)
            if valor not in (None, ""):
                setattr(persona, campo, valor)
        persona.estado = b.estado
        persona.activo = b.estado == "Activo"
    session.flush()
    return resumen


def importar_bomberos_desde_excel(ruta: Path) -> Tuple[ResumenImportacion, Optional[int]]:
    """Aplica el 'Reporte de bomberos' a la base (una transacción) y
    actualiza la copia local del padrón que usan los selectores de RUBA.
    Devuelve (resumen, cantidad de activos en el padrón | None si no se
    pudo actualizar el padrón -- el motivo va en los avisos)."""
    from app.core.catalogos import importar_padron
    from app.db import get_session

    lectura = leer_reporte_bomberos(ruta)
    with get_session() as session:
        resumen = upsert_bomberos(session, lectura.bomberos)
    resumen.ignoradas += lectura.ignoradas
    resumen.avisos = lectura.avisos + resumen.avisos
    try:
        activos: Optional[int] = len(importar_padron(Path(ruta)))
    except (OSError, ValueError) as e:
        activos = None
        resumen.avisos.append(f"La base se actualizó, pero no el padrón de los selectores de RUBA: {e}")
    log.info("Padrón %s: %s", Path(ruta).name, resumen.texto_personal())
    return resumen, activos
